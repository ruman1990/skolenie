

import csv
from decimal import Decimal
with open("data.csv","r",encoding="utf-8") as f:
    reader = csv.reader(f)
    header = reader.__next__()
    data = [[x[0],x[1],x[2],Decimal(x[3])] for x in reader]

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.chart import BarChart, Reference

wb = Workbook()
ws = wb.active
ws.title = "Platy"

ws.append(header[1:])

for x in data:
    ws.append(x[1:])

ws["A1"].font = Font(bold=True, color="0000FF")
ws["B1"].font = Font(bold=True, color="0000FF")
ws["C1"].font = Font(bold=True, color="0000FF")

last_row = ws.max_row

ws.append([None,None,f"=AVERAGE(C2:C{last_row})"])


for r in ws.iter_rows(min_row=2):
    for x in r:
        x.number_format = "#,##0.00"

# 3. Vytvorenie grafu
graf = BarChart()
graf.title = "Platy zamestnancov"
graf.y_axis.title = "€"
graf.x_axis.title = "Meno"

# 4. Určenie dát pre graf (iba stĺpec „Spolu €“)
data = Reference(ws, min_col=3, min_row=2,max_row=last_row)
kategorie = Reference(ws, min_col=2, min_row=2, max_row=last_row)

graf.add_data(data, titles_from_data=True)
graf.set_categories(kategorie)

# 5. Vloženie grafu
ws.add_chart(graf, "F2")

wb.save("vystup.xlsx")